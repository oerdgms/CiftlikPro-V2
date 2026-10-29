"""Animal reports must agree with the active herd and preserve archive status."""
import io
import tempfile
import threading
import unittest
import urllib.request
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlencode

from openpyxl import load_workbook
from pypdf import PdfReader

from test_reports_import import server
import mobile_reference


class AnimalReportMembershipTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_patch = patch.object(server, 'DB', Path(self.temp.name) / 'reports.db')
        self.db_patch.start()
        server.init_db()
        self.day = server.date.today().isoformat()
        self.adults = []
        with server.db() as c:
            for index in range(55):
                self.adults.append(c.execute('''insert into animals
                    (tag,nickname,gender,status,paddock,breed,birth_date,purchase_date)
                    values(?,?,?,'Aktif','SA1','Simental','2020-01-01','2025-01-01')''',
                    (f'TR583010{index:03d}', 'İnci Sağmal' if index == 0 else f'Hayvan {index}',
                     'Dişi' if index < 30 else 'Erkek')).lastrowid)
            for index in range(12):
                c.execute('''insert into calves(tag,mother_id,gender,status,paddock,birth_date)
                    values(?,?,'Dişi','Aktif','BUZ',?)''',
                    (f'TR583020{index:03d}', self.adults[0], self.day))
            self.sold_tags = {f'TR583030{index:03d}' for index in range(4)}
            for tag in sorted(self.sold_tags):
                c.execute('''insert into calves(tag,mother_id,gender,status,paddock,birth_date)
                    values(?,?,'Erkek','Satıldı','BUZ',?)''', (tag, self.adults[0], self.day))
            c.execute("insert into animals(tag,gender,status) values('TR583040000','Dişi','Satıldı')")
            self.sold_tags.add('TR583040000')

    def tearDown(self):
        self.db_patch.stop()
        self.temp.cleanup()

    def test_active_67_matches_herd_and_excludes_four_sold_calves(self):
        rows = server.animal_report_rows()
        with server.db() as c:
            herd = mobile_reference.herd_records(vars(server), c)
        self.assertEqual(len(rows), 67)
        self.assertEqual({r['tag'] for r in rows}, {r['tag'] for r in herd})
        self.assertFalse(self.sold_tags & {r['tag'] for r in rows})
        self.assertEqual({r['status'] for r in rows}, {'Aktif'})

    def test_sold_and_all_filters_keep_adults_and_calves_with_real_status(self):
        sold = server.animal_report_rows(status='Satıldı')
        self.assertEqual({r['tag'] for r in sold}, self.sold_tags)
        self.assertEqual({r['status'] for r in sold}, {'Satıldı'})
        self.assertEqual(len(server.animal_report_rows(status='Tümü')), 72)
        self.assertEqual(len(server.animal_report_rows(group='calves', status='Satıldı')), 4)
        self.assertEqual(len(server.animal_report_rows(group='calves', status='Tümü')), 16)
        with server.db() as c:
            c.execute("update calves set status='Kesildi' where tag='TR583030000'")
        cut = server.animal_report_rows(group='calves', status='Kesildi')
        self.assertEqual([(r['tag'], r['status']) for r in cut], [('TR583030000', 'Kesildi')])

    def test_promoted_calf_is_counted_once_and_mother_is_preserved(self):
        with server.db() as c:
            c.execute('''insert into calves(tag,mother_id,gender,status,promoted_animal_id,birth_date)
                values('TR583010001',?,'Dişi','Aktif',?,?)''', (self.adults[0], self.adults[1], self.day))
        rows = server.animal_report_rows()
        self.assertEqual(len(rows), 67)
        promoted = [r for r in rows if r['tag'] == 'TR583010001']
        self.assertEqual(len(promoted), 1)
        self.assertEqual(promoted[0]['group'], 'Dişi')
        self.assertEqual(promoted[0]['mother_tag'], 'TR583010000')

    def test_losses_are_excluded_even_with_stale_active_status(self):
        with server.db() as c:
            for source, column in [('animal', 'animal_id'), ('calf', 'calf_id')]:
                tag = 'TR583050000' if source == 'animal' else 'TR583050001'
                if source == 'animal':
                    sid = c.execute("insert into animals(tag,gender,status) values(?,'Dişi','Aktif')", (tag,)).lastrowid
                else:
                    sid = c.execute("insert into calves(tag,mother_id,birth_date,gender,status) values(?,?,?,'Dişi','Aktif')", (tag, self.adults[0], self.day)).lastrowid
                c.execute(f'''insert into animal_losses(subject_type,{column},event_date,event_type,created_at)
                    values(?,?,?,'Öldü',?)''', (source, sid, self.day, self.day))
        self.assertEqual(len(server.animal_report_rows()), 67)
        lost = server.animal_report_rows(status='Öldü')
        self.assertEqual({r['tag'] for r in lost}, {'TR583050000', 'TR583050001'})
        self.assertEqual({r['status'] for r in lost}, {'Öldü'})

    def test_group_search_pad_and_null_status_preserve_identity_fields(self):
        self.assertEqual(len(server.animal_report_rows(group='female')), 30)
        self.assertEqual(len(server.animal_report_rows(group='male')), 25)
        self.assertEqual(len(server.animal_report_rows(group='calves')), 12)
        rows = server.animal_report_rows(search='INCI SAGMAL', paddock='SA1')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['arrival_date'], '2025-01-01')
        self.assertEqual(rows[0]['nickname'], 'İnci Sağmal')
        self.assertEqual(server.animal_report_rows(search='INCI SAGMAL', paddock='BUZ'), [])
        calf = server.animal_report_rows(search='TR583020000')[0]
        self.assertEqual(calf['mother_tag'], 'TR583010000')
        with server.db() as c:
            c.execute("update animals set status=null where id=?", (self.adults[0],))
            c.execute("update calves set status='' where tag='TR583020000'")
        self.assertEqual(len(server.animal_report_rows()), 67)

    def test_report_selection_does_not_change_source_records(self):
        def snapshot():
            with server.db() as c:
                return {table: [tuple(r) for r in c.execute(f'select * from {table} order by id')]
                        for table in ('animals', 'calves', 'finance', 'animal_losses')}
        before = snapshot()
        for status in ('Aktif', 'Satıldı', 'Kesildi', 'Tümü'):
            server.animal_report_rows(status=status)
        self.assertEqual(snapshot(), before)

    def test_http_screen_preview_pdf_and_excel_use_the_same_membership(self):
        http = server.QuietThreadingHTTPServer(('127.0.0.1', 0), server.App)
        threading.Thread(target=http.serve_forever, daemon=True).start()
        server.SESSIONS['report-membership'] = {'username': 'admin', 'role': 'admin'}
        base = 'http://127.0.0.1:' + str(http.server_port)
        def get(path):
            request = urllib.request.Request(base + path, headers={'Cookie': 'sid=report-membership'})
            with urllib.request.urlopen(request, timeout=20) as response:
                return response.read()
        try:
            with patch.object(server, 'license_status', return_value=(True, {}, '')):
                for status, count in [('Aktif', 67), ('Satıldı', 5)]:
                    query = urlencode([('tab', 'animals'), ('animal_status', status),
                                       ('columns_mode', 'custom'), ('columns', 'tag'), ('columns', 'status')])
                    screen = get('/reports?' + query).decode()
                    preview = get('/reports/animals/print?' + query).decode()
                    self.assertIn(f'{count} kayıt', screen)
                    self.assertIn(f'Toplam Hayvan: {count}', preview)
                    for tag in self.sold_tags:
                        self.assertEqual(tag in screen, status == 'Satıldı')
                        self.assertEqual(tag in preview, status == 'Satıldı')
                    workbook = load_workbook(io.BytesIO(get('/reports/animals.xlsx?' + query)), read_only=True)
                    try:
                        excel = list(workbook.active.iter_rows(min_row=5, values_only=True))
                        self.assertEqual(len(excel), count)
                        self.assertEqual({r[1] for r in excel}, {status})
                        self.assertEqual(bool(self.sold_tags & {r[0] for r in excel}), status == 'Satıldı')
                    finally:
                        workbook.close()
                    pdf = PdfReader(io.BytesIO(get('/reports/animals.pdf?' + query)))
                    text = '\n'.join(page.extract_text() for page in pdf.pages)
                    for tag in self.sold_tags:
                        self.assertEqual(tag in text, status == 'Satıldı')
        finally:
            http.shutdown()
            http.server_close()
            server.SESSIONS.pop('report-membership', None)


if __name__ == '__main__':
    unittest.main()
